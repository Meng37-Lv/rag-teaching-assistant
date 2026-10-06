from __future__ import annotations
import hashlib,io,secrets,sqlite3,unicodedata,os
from io import BytesIO
from urllib.parse import urlparse
from fastapi.responses import StreamingResponse
from datetime import datetime,timezone
from uuid import uuid4
from fastapi import APIRouter,File,HTTPException,Request,UploadFile
from pydantic import BaseModel,ConfigDict,Field
from src.auth import get_accounts
def now(): return datetime.now(timezone.utc).isoformat()
def h(x): return hashlib.sha256(x.encode()).hexdigest()
def clean_name(x): return unicodedata.normalize('NFC',''.join(' ' if unicodedata.category(c)=='Zs' else c for c in str(x))).strip()
def db(r): return get_accounts(r).db_path
def own(r,c):
 if not get_accounts(r).owns(r.state.account['account_id'],c): raise HTTPException(404,'课程不存在或无权访问。')
def con(p):
 c=sqlite3.connect(p); c.row_factory=sqlite3.Row; c.execute('PRAGMA foreign_keys=ON'); return c
class ClassCreate(BaseModel):
 model_config=ConfigDict(extra='forbid'); name:str=Field(min_length=1,max_length=100)
class StudentJoin(BaseModel): name:str=Field(min_length=1,max_length=100)
class MessageCreate(BaseModel): body:str=Field(min_length=1,max_length=1000)
router=APIRouter(prefix='/api',tags=['classrooms'])
@router.post('/courses/{course_id}/classes')
def cc(course_id,payload:ClassCreate,request:Request):
 own(request,course_id); n=clean_name(payload.name); cid=str(uuid4()); t=now(); c=con(db(request))
 try:
  with c:c.execute('INSERT INTO classes VALUES (?,?,?,?,?)',(cid,course_id,n,t,t))
 except sqlite3.IntegrityError: raise HTTPException(409,'同课程下班级名称已存在。')
 finally:c.close()
 return {'class_id':cid,'course_id':course_id,'name':n,'created_at':t,'updated_at':t}
@router.get('/courses/{course_id}/classes')
def lc(course_id,request:Request):
 own(request,course_id); c=con(db(request)); x=c.execute('SELECT * FROM classes WHERE course_id=?',(course_id,)).fetchall(); c.close(); return [dict(r) for r in x]
@router.post('/classes/{class_id}/import')
async def imp(class_id,request:Request,file:UploadFile=File(...)):
 if not(file.filename or '').lower().endswith('.xlsx'):raise HTTPException(400,'仅支持.xlsx文件。')
 try:
  from openpyxl import load_workbook
  rows=list(load_workbook(filename=io.BytesIO(await file.read()),read_only=True,data_only=True).active.iter_rows(values_only=True))
 except Exception:raise HTTPException(400,'Excel解析失败。')
 if not rows or list(rows[0])[:2]!=['姓名','学号']:raise HTTPException(400,'首行必须为姓名、学号。')
 vals=[]; seen=set()
 for i,row in enumerate(rows[1:],2):
  n,no=clean_name(row[0] or ''),str(row[1] or '').strip()
  if not n or not no:raise HTTPException(400,f'第{i}行姓名和学号不能为空。')
  if no in seen:raise HTTPException(400,f'第{i}行学号重复或姓名冲突。')
  seen.add(no); vals.append((n,no))
 c=con(db(request))
 try:
  row=c.execute('SELECT course_id FROM classes WHERE class_id=?',(class_id,)).fetchone()
  if not row:raise HTTPException(404,'班级不存在。')
  own(request,row['course_id']); old={r[0]:r[1] for r in c.execute('SELECT student_no,name FROM class_students WHERE class_id=?',(class_id,))}
  with c:
   for n,no in vals:
    if no in old:
     if old[no] != n: raise HTTPException(400,f'学号{no}已有不同姓名，无法增量导入。')
    else: c.execute('INSERT INTO class_students VALUES (?,?,?,?,?)',(str(uuid4()),class_id,n,no,now()))
 finally:c.close()
 return {'added':sum(no not in old for _,no in vals),'total_rows':len(vals)}
@router.post('/classes/{class_id}/sessions')
def ss(class_id,request:Request):
 c=con(db(request)); row=c.execute('SELECT course_id FROM classes WHERE class_id=?',(class_id,)).fetchone()
 if not row:c.close();raise HTTPException(404,'班级不存在。')
 own(request,row['course_id']); token,sid=secrets.token_urlsafe(32),str(uuid4())
 try:
  with c:
   c.execute('INSERT INTO classroom_sessions VALUES (?,?,?,?,?,?)',(sid,class_id,h(token),'active',now(),None))
 except sqlite3.IntegrityError:raise HTTPException(409,'该班级已有进行中的课堂。')
 finally:c.close()
 return {'session_id':sid,'join_token':token,'status':'active'}
@router.post('/sessions/{session_id}/end')
def es(session_id,request:Request):
 c=con(db(request)); row=c.execute('SELECT c.course_id,c.class_id FROM classroom_sessions s JOIN classes c ON c.class_id=s.class_id WHERE s.session_id=?',(session_id,)).fetchone()
 if not row:c.close();raise HTTPException(404,'课堂不存在。')
 own(request,row['course_id']); c.execute("UPDATE classroom_sessions SET status='ended',ended_at=? WHERE session_id=?",(now(),session_id));c.commit();c.close();return {'status':'ended'}
@router.post('/join/{join_token}')
def join(join_token,payload:StudentJoin,request:Request):
 n=clean_name(payload.name); c=con(db(request)); row=c.execute("SELECT session_id,class_id FROM classroom_sessions WHERE join_token_hash=? AND status='active'",(h(join_token),)).fetchone()
 if not row:c.close();raise HTTPException(404,'课堂不存在或已结束。')
 m=c.execute('SELECT * FROM class_students WHERE class_id=? AND name=?',(row['class_id'],n)).fetchall(); st='confirmed' if len(m)==1 else 'pending'; sid,tok=str(uuid4()),secrets.token_urlsafe(32); student=m[0]['student_id'] if len(m)==1 else None
 c.execute('INSERT INTO classroom_students (session_student_id,session_id,student_id,student_token_hash,status,joined_at,requested_name) VALUES (?,?,?,?,?,?,?)',(sid,row['session_id'],student,h(tok),st,now(),n));c.commit();c.close();return {'student_token':tok,'status':st}
@router.post('/sessions/{session_id}/messages')
def pm(session_id,payload:MessageCreate,request:Request):
 body=payload.body.strip(); token=request.headers.get('X-Student-Session','');c=con(db(request)); row=c.execute("SELECT session_student_id FROM classroom_students WHERE session_id=? AND student_token_hash=? AND status='confirmed'",(session_id,h(token))).fetchone();active=c.execute("SELECT 1 FROM classroom_sessions WHERE session_id=? AND status='active'",(session_id,)).fetchone()
 if not row or not active:c.close();raise HTTPException(403,'学生会话无效或课堂已结束。')
 mid=str(uuid4());c.execute('INSERT INTO classroom_messages VALUES (?,?,?,?,?)',(mid,session_id,row['session_student_id'],body,now()));c.commit();c.close();return {'message_id':mid}
@router.get('/sessions/{session_id}/messages')
def lm(session_id,request:Request):
 c=con(db(request));row=c.execute('SELECT c.course_id,c.class_id FROM classroom_sessions s JOIN classes c ON c.class_id=s.class_id WHERE s.session_id=?',(session_id,)).fetchone()
 if not row:c.close();raise HTTPException(404,'课堂不存在。')
 own(request,row['course_id']);rows=c.execute('SELECT m.message_id,m.body,m.created_at,cs.name FROM classroom_messages m JOIN classroom_students ss ON ss.session_student_id=m.session_student_id JOIN class_students cs ON cs.student_id=ss.student_id WHERE m.session_id=? ORDER BY m.created_at',(session_id,)).fetchall();c.close();return [dict(r) for r in rows]






@router.get('/sessions/{session_id}/pending')
def pending(session_id, request: Request):
 c=con(db(request)); row=c.execute('SELECT c.course_id FROM classroom_sessions s JOIN classes c ON c.class_id=s.class_id WHERE s.session_id=?',(session_id,)).fetchone()
 if not row: c.close(); raise HTTPException(404,'课堂不存在。')
 own(request,row['course_id']); rows=c.execute("SELECT cs.session_student_id,cs.status,cs.joined_at,cs.requested_name,s.class_id FROM classroom_students cs JOIN classroom_sessions s ON s.session_id=cs.session_id WHERE cs.session_id=? AND cs.status='pending'",(session_id,)).fetchall(); out=[]
 for item in rows:
  value=dict(item); value['candidates']=[dict(x) for x in c.execute('SELECT student_id,name,student_no FROM class_students WHERE class_id=? AND name=?',(item['class_id'],item['requested_name'])).fetchall()]; value.pop('class_id',None); out.append(value)
 c.close(); return out

@router.post('/sessions/{session_id}/students/{session_student_id}/confirm')
def confirm(session_id, session_student_id, request: Request, student_id: str):
 c=con(db(request)); row=c.execute('SELECT c.course_id,c.class_id FROM classroom_sessions s JOIN classes c ON c.class_id=s.class_id WHERE s.session_id=?',(session_id,)).fetchone()
 if not row: c.close(); raise HTTPException(404,'课堂不存在。')
 own(request,row['course_id']); candidate=c.execute('SELECT class_id FROM class_students WHERE student_id=?',(student_id,)).fetchone(); link=c.execute('SELECT session_student_id FROM classroom_students WHERE session_student_id=? AND session_id=?',(session_student_id,session_id)).fetchone()
 if not candidate or not link or candidate['class_id'] != row['class_id']: c.close(); raise HTTPException(404,'学生或申请不存在。')
 c.execute("UPDATE classroom_students SET student_id=?,status='confirmed' WHERE session_student_id=?",(student_id,session_student_id)); c.commit(); c.close(); c2=con(db(request)); out=c2.execute('SELECT st.name,st.student_no FROM class_students st WHERE st.student_id=?',(student_id,)).fetchone(); c2.close(); return {'status':'confirmed','name':out['name'],'student_no':out['student_no']}

@router.post('/sessions/{session_id}/students/{session_student_id}/reject')
def reject(session_id, session_student_id, request: Request):
 c=con(db(request)); row=c.execute('SELECT c.course_id FROM classroom_sessions s JOIN classes c ON c.class_id=s.class_id WHERE s.session_id=?',(session_id,)).fetchone()
 if not row: c.close(); raise HTTPException(404,'课堂不存在。')
 own(request,row['course_id']); cur=c.execute("UPDATE classroom_students SET status='rejected' WHERE session_student_id=? AND session_id=? AND status='pending'",(session_student_id,session_id)); c.commit(); c.close()
 if not cur.rowcount: raise HTTPException(404,'待确认申请不存在。')
 return {'status':'rejected'}

@router.get('/sessions/{session_id}/student-status')
def student_status(session_id, request: Request):
 token=request.headers.get('X-Student-Session',''); c=con(db(request)); row=c.execute("SELECT cs.status,cs.requested_name,s.status AS session_status FROM classroom_students cs JOIN classroom_sessions s ON s.session_id=cs.session_id WHERE cs.session_id=? AND cs.student_token_hash=?",(session_id,h(token))).fetchone()
 if not row: c.close(); raise HTTPException(404,'学生会话不存在。')
 msgs=c.execute('SELECT body,created_at FROM classroom_messages WHERE session_id=? AND session_student_id=(SELECT session_student_id FROM classroom_students WHERE session_id=? AND student_token_hash=?) ORDER BY created_at',(session_id,session_id,h(token))).fetchall(); c.close(); return {'status':row['status'],'session_status':row['session_status'],'name':row['requested_name'],'messages':[dict(x) for x in msgs]}

@router.delete('/classes/{class_id}')
def delete_class(class_id, request: Request):
 c=con(db(request)); row=c.execute('SELECT course_id FROM classes WHERE class_id=?',(class_id,)).fetchone()
 if not row: c.close(); raise HTTPException(404,'班级不存在。')
 own(request,row['course_id'])
 if c.execute('SELECT 1 FROM class_students WHERE class_id=? UNION SELECT 1 FROM classroom_sessions WHERE class_id=?',(class_id,class_id)).fetchone(): c.close(); raise HTTPException(409,'班级已有名单或课堂记录，不能删除。')
 c.execute('DELETE FROM classes WHERE class_id=?',(class_id,)); c.commit(); c.close(); return None

@router.get('/student/state')
def student_state(request: Request):
 token=request.headers.get('X-Student-Session',''); c=con(db(request)); row=c.execute('SELECT cs.status,cs.requested_name,s.status AS session_status,cs.session_id FROM classroom_students cs JOIN classroom_sessions s ON s.session_id=cs.session_id WHERE cs.student_token_hash=?',(h(token),)).fetchone()
 if not row: c.close(); raise HTTPException(404,'学生会话不存在。')
 messages=c.execute('SELECT body,created_at FROM classroom_messages WHERE session_student_id=(SELECT session_student_id FROM classroom_students WHERE student_token_hash=?) ORDER BY created_at',(h(token),)).fetchall(); c.close(); return {'status':row['status'],'session_status':row['session_status'],'name':row['requested_name'],'messages':[dict(x) for x in messages]}
@router.post('/student/messages')
def student_message(payload: MessageCreate, request: Request):
 body=payload.body.strip(); token=request.headers.get('X-Student-Session',''); c=con(db(request)); row=c.execute("SELECT cs.session_student_id FROM classroom_students cs JOIN classroom_sessions s ON s.session_id=cs.session_id WHERE cs.student_token_hash=? AND cs.status='confirmed' AND s.status='active'",(h(token),)).fetchone()
 if not row: c.close(); raise HTTPException(403,'学生会话无效或课堂已结束。')
 mid=str(uuid4()); c.execute('INSERT INTO classroom_messages VALUES (?,?,?,?,?)',(mid, c.execute('SELECT session_id FROM classroom_students WHERE session_student_id=?',(row['session_student_id'],)).fetchone()[0], row['session_student_id'], body, now())); c.commit(); c.close(); return {'message_id':mid}


def _teacher_session_row(request, session_id):
 c=con(db(request)); row=c.execute('SELECT s.session_id,s.status,s.join_token_hash,c.class_id,c.course_id FROM classroom_sessions s JOIN classes c ON c.class_id=s.class_id WHERE s.session_id=?',(session_id,)).fetchone(); return c,row

def _public_join_url(token: str) -> str:
 base=os.getenv('PUBLIC_WEB_BASE_URL','').strip().rstrip('/')
 parsed=urlparse(base)
 if parsed.scheme != 'https' or not parsed.hostname or parsed.hostname.lower() in {'localhost','127.0.0.1','::1'}:
  raise HTTPException(503,'请先配置公网HTTPS地址。')
 return f'{base}/?join={token}'

def _join_token_header(request: Request) -> str:
 token=request.headers.get('X-Join-Token','').strip()
 if not token: raise HTTPException(400,'缺少课堂加入令牌。')
 return token

def _join_info_for_teacher(session_id, request: Request):
 token=_join_token_header(request)
 c,row=_teacher_session_row(request,session_id)
 if not row: c.close(); raise HTTPException(404,'课堂不存在。')
 own(request,row['course_id'])
 if row['status']!='active': c.close(); raise HTTPException(404,'课堂已结束。')
 if h(token) != row['join_token_hash']: c.close(); raise HTTPException(404,'课堂加入令牌无效。')
 url=_public_join_url(token); c.close(); return {'join_url':url,'qr_available':True}

@router.get('/sessions/{session_id}/join-info')
def join_info(session_id, request: Request):
 return _join_info_for_teacher(session_id,request)

@router.get('/sessions/{session_id}/qrcode')
def qrcode_png(session_id, request: Request):
 info=_join_info_for_teacher(session_id,request)
 try:
  import qrcode
  image=qrcode.make(info['join_url']); stream=BytesIO(); image.save(stream,format='PNG'); stream.seek(0)
 except HTTPException: raise
 except Exception: raise HTTPException(500,'二维码生成失败。') from None
 return StreamingResponse(stream,media_type='image/png',headers={'Cache-Control':'no-store'})
