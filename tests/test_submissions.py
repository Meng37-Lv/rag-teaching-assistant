from __future__ import annotations

import io
import json
import os
import tempfile
import unittest
import gc
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook
from pptx import Presentation

from src.course_data import CourseStore, get_course_store
from src.web_api import app
from tests.auth_helpers import login_test_client
from src.config import Settings
from src.llm_client import LLMCompletion
from src.material_parser import MaterialSegment, ParsedMaterial
from src.report_question_service import ReportQuestionService
from tests.test_report_question_service import grounded_output
from src.preview_converter import PreviewConversionError, PreviewResult


class SubmissionApiTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.old_data_dir = os.environ.get("RAG_APP_DATA_DIR")
        os.environ["RAG_APP_DATA_DIR"] = self.tmp.name
        self.store = CourseStore(Path(self.tmp.name) / "db.sqlite")
        app.dependency_overrides[get_course_store] = lambda: self.store
        self.client = TestClient(app)
        login_test_client(self.client, self.store)
        course = self.client.post("/api/courses", json={"name": "提交测试课程"}).json()
        classroom = self.client.post(f"/api/courses/{course['id']}/classes", json={"name": "一班"}).json()
        workbook = Workbook(); sheet = workbook.active
        sheet.append(["姓名", "学号"]); sheet.append(["张三", "001"])
        content = io.BytesIO(); workbook.save(content); content.seek(0)
        self.client.post(f"/api/classes/{classroom['class_id']}/import", files={"file": ("roster.xlsx", content.getvalue())})
        self.session = self.client.post(f"/api/classes/{classroom['class_id']}/sessions").json()
        self.class_id = classroom["class_id"]
        self.teacher_token = self.client.headers["Authorization"]
        student = self.client.post(f"/api/join/{self.session['join_token']}", json={"name": "张三"}).json()
        self.student_token = student["student_token"]
        self.llm_messages = []

        owner = self
        class CapturingLLM:
            def complete(inner, messages, max_tokens):
                owner.llm_messages.append(messages)
                payload = json.loads(messages[1]["content"].split("<material>\n")[1].split("\n</material>")[0])
                segments = tuple(MaterialSegment(**item) for item in payload["片段"])
                material = ParsedMaterial(payload["材料名称"], payload["类型"], "\n".join(item.text for item in segments), segments)
                return LLMCompletion(json.dumps(grounded_output(material), ensure_ascii=False), "stop", 100, max_tokens)

        settings = Settings("test", "https://example.invalid", "test", "high", False, 12, 3, 1800, 600)
        self.service_patch = patch("src.submissions._get_submission_question_service", return_value=ReportQuestionService(settings, CapturingLLM()))
        self.service_patch.start()

    def tearDown(self):
        self.service_patch.stop(); self.client.close(); app.dependency_overrides.clear(); gc.collect(); self.tmp.cleanup()
        if self.old_data_dir is None: os.environ.pop("RAG_APP_DATA_DIR", None)
        else: os.environ["RAG_APP_DATA_DIR"] = self.old_data_dir

    def _student_headers(self):
        return {"X-Student-Session": self.student_token}

    def test_submit_list_detail_preview_and_question(self):
        upload = self.client.post(
            f"/api/sessions/{self.session['session_id']}/submissions",
            headers=self._student_headers(),
            files=[("files", ("report.md", "模型方案采用对照实验验证结果。\n当前实验仅覆盖室内场景。".encode(), "text/markdown")), ("files", ("notes.txt", b"plain text", "text/plain"))],
        )
        self.assertEqual(upload.status_code, 200)
        submission = upload.json(); self.assertEqual(len(submission["files"]), 2)
        listing = self.client.get(f"/api/sessions/{self.session['session_id']}/submissions")
        self.assertEqual(listing.status_code, 200); self.assertEqual(listing.json()["total"], 1)
        self.assertEqual(self.client.get(f"/api/sessions/{self.session['session_id']}/submissions?student_no=001").json()["total"], 1)
        detail = self.client.get(f"/api/submissions/{submission['submission_id']}"); self.assertEqual(detail.status_code, 200)
        file_id = next(item["file_id"] for item in submission["files"] if item["file_type"] == "md")
        preview = self.client.get(f"/api/submission-files/{file_id}/preview"); self.assertEqual(preview.status_code, 200); self.assertEqual(preview.json()["preview_type"], "markdown")
        self.assertEqual(self.client.get(f"/api/submission-files/{file_id}/content").status_code, 200)
        question = self.client.post(f"/api/submission-files/{file_id}/generate-question")
        self.assertEqual(question.status_code, 200); self.assertTrue(question.json()["response"]["citations"]); self.assertIn("模型方案", json.dumps(self.llm_messages, ensure_ascii=False))

        import sqlite3
        with closing(sqlite3.connect(self.store.db_path)) as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM teaching_questions").fetchone()[0], 1)

    def test_invalid_llm_output_does_not_save_question(self):
        from src.report_question_service import ReportQuestionOutputError
        upload = self.client.post(f"/api/sessions/{self.session['session_id']}/submissions", headers=self._student_headers(), files={"files": ("report.txt", "模型方案和实验结果".encode(), "text/plain")})
        file_id = upload.json()["files"][0]["file_id"]
        with patch("src.submissions._get_submission_question_service", side_effect=ReportQuestionOutputError("invalid", "{}")):
            failed = self.client.post(f"/api/submission-files/{file_id}/generate-question")
        self.assertEqual(failed.status_code, 502); self.assertEqual(failed.json()["detail"]["code"], "LLM_OUTPUT_INVALID")
        import sqlite3
        with closing(sqlite3.connect(self.store.db_path)) as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM teaching_questions").fetchone()[0], 0)

    def test_office_preview_endpoint_returns_pdf_and_failure_is_explicit(self):
        presentation = Presentation(); presentation.slides.add_slide(presentation.slide_layouts[6])
        content = io.BytesIO(); presentation.save(content); content.seek(0)
        upload = self.client.post(f"/api/sessions/{self.session['session_id']}/submissions", headers=self._student_headers(), files={"files": ("slides.pptx", content.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
        file_id = upload.json()["files"][0]["file_id"]
        preview_pdf = Path(self.tmp.name) / "converted.pdf"; preview_pdf.write_bytes(b"%PDF-1.7\nconverted")
        result = PreviewResult("PREVIEW_OK", preview_pdf, "cache-key", "lo-test", False)
        self.client.headers.pop("Authorization", None)
        with patch("src.submissions.convert_to_pdf", return_value=result), patch("src.submissions.get_submission_preview_dir", return_value=preview_pdf.parent):
            metadata = self.client.get(f"/api/submission-files/{file_id}/preview", headers=self._student_headers())
            self.assertEqual(metadata.status_code, 200)
            self.assertEqual(metadata.json()["preview_type"], "pdf")
            rendered = self.client.get(f"/api/submission-files/{file_id}/preview/pdf", headers=self._student_headers())
            self.assertEqual(rendered.status_code, 200)
            self.assertTrue(rendered.content.startswith(b"%PDF-"))

        self.client.headers["Authorization"] = self.teacher_token
        with patch("src.submissions.convert_to_pdf", side_effect=PreviewConversionError("PREVIEW_TIMEOUT", "预览转换超时", status_code=504)), patch("src.submissions.get_submission_preview_dir", return_value=preview_pdf.parent):
            failed = self.client.get(f"/api/submission-files/{file_id}/preview")
        self.assertEqual(failed.status_code, 504)
        self.assertEqual(failed.json()["detail"]["code"], "PREVIEW_TIMEOUT")
        detail = self.client.get(f"/api/submissions/{upload.json()['submission_id']}")
        self.assertEqual(detail.json()["status"], "submitted")

    def test_submission_validation_and_classroom_state(self):
        bad = self.client.post(f"/api/sessions/{self.session['session_id']}/submissions", headers=self._student_headers(), files={"files": ("bad.exe", b"x", "application/octet-stream")})
        self.assertEqual(bad.status_code, 400); self.assertEqual(bad.json()["detail"]["code"], "FILE_TYPE_UNSUPPORTED")
        self.assertEqual(self.client.post(f"/api/sessions/{self.session['session_id']}/submissions", headers={"X-Student-Session": "invalid"}, files={"files": ("a.txt", b"x", "text/plain")}).status_code, 403)
        valid = self.client.post(f"/api/sessions/{self.session['session_id']}/submissions", headers=self._student_headers(), files={"files": ("a.txt", b"x", "text/plain")}).json()
        self.client.post(f"/api/sessions/{self.session['session_id']}/end")
        ended = self.client.post(f"/api/sessions/{self.session['session_id']}/submissions", headers=self._student_headers(), files={"files": ("a.txt", b"x", "text/plain")})
        self.assertEqual(ended.status_code, 409); self.assertEqual(ended.json()["detail"]["code"], "CLASSROOM_ENDED")
        self.client.headers.pop("X-Student-Session", None)
        self.assertEqual(self.client.post(f"/api/submission-files/{valid['files'][0]['file_id']}/generate-question").status_code, 409)

    def test_cross_teacher_student_and_storage_boundaries(self):
        uploaded = self.client.post(f"/api/sessions/{self.session['session_id']}/submissions", headers=self._student_headers(), files={"files": ("..\\evil.txt", b"<script>alert(1)</script>", "text/plain")})
        self.assertEqual(uploaded.status_code, 200)
        submission = uploaded.json(); file_id = submission["files"][0]["file_id"]
        self.assertEqual(submission["files"][0]["filename"], "evil.txt")
        # A separate teacher account cannot traverse the first teacher's class.
        registration = self.client.post("/api/auth/register", json={"account_name": "教师B", "phone": "13900000002", "password": "fixture", "confirm_password": "fixture"})
        self.assertEqual(registration.status_code, 201)
        other_login = self.client.post("/api/auth/login", json={"phone": "13900000002", "password": "fixture"}).json()
        self.client.headers["Authorization"] = "Bearer " + other_login["token"]
        self.assertEqual(self.client.get(f"/api/sessions/{self.session['session_id']}/submissions").status_code, 404)
        self.assertEqual(self.client.get(f"/api/submissions/{submission['submission_id']}").status_code, 404)
        self.assertEqual(self.client.get(f"/api/submission-files/{file_id}/preview").status_code, 404)
        self.assertEqual(self.client.get(f"/api/submission-files/{file_id}/content").status_code, 404)
        self.assertEqual(self.client.post(f"/api/submission-files/{file_id}/generate-question").status_code, 404)

        # The original student token remains the sole owner of the submission.
        self.client.headers.pop("Authorization")
        self.client.headers["X-Student-Session"] = self.student_token
        self.assertEqual(self.client.get(f"/api/submission-files/{file_id}/content").status_code, 200)

        # A second confirmed student in the same class has a separate scope.
        self.client.headers["Authorization"] = self.teacher_token
        roster = Workbook(); sheet = roster.active; sheet.append(["姓名", "学号"]); sheet.append(["张三", "001"]); sheet.append(["李四", "002"])
        roster_bytes = io.BytesIO(); roster.save(roster_bytes); roster_bytes.seek(0)
        self.assertEqual(self.client.post(f"/api/classes/{self.class_id}/import", files={"file": ("roster.xlsx", roster_bytes.getvalue())}).status_code, 200)
        student_b = self.client.post(f"/api/join/{self.session['join_token']}", json={"name": "李四"}).json()
        second = self.client.post(f"/api/sessions/{self.session['session_id']}/submissions", headers={"X-Student-Session": student_b["student_token"]}, files={"files": ("b.txt", b"student-b", "text/plain")}).json()
        self.client.headers.pop("Authorization")
        self.client.headers["X-Student-Session"] = self.student_token
        self.assertEqual(self.client.get(f"/api/sessions/{self.session['session_id']}/submissions").json()["total"], 1)
        self.assertEqual(self.client.get(f"/api/submissions/{second['submission_id']}").status_code, 403)
        self.assertEqual(self.client.get(f"/api/submission-files/{second['files'][0]['file_id']}/preview").status_code, 403)
        self.assertEqual(self.client.get(f"/api/submission-files/{second['files'][0]['file_id']}/content").status_code, 403)
        self.client.headers["X-Student-Session"] = "invalid"

        # A pending student cannot upload before confirmation.
        pending = self.client.post(f"/api/join/{self.session['join_token']}", json={"name": "未在名单学生"}).json()
        self.assertEqual(pending["status"], "pending")
        self.assertEqual(self.client.post(f"/api/sessions/{self.session['session_id']}/submissions", headers={"X-Student-Session": pending["student_token"]}, files={"files": ("a.txt", b"x", "text/plain")}).status_code, 403)

        # Corrupt binary files fail before a submission row or storage path is committed.
        self.assertEqual(self.client.post(f"/api/sessions/{self.session['session_id']}/submissions", headers=self._student_headers(), files={"files": ("broken.pptx", b"not-a-pptx", "application/octet-stream")}).status_code, 400)
        import sqlite3
        with closing(sqlite3.connect(self.store.db_path)) as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM submissions").fetchone()[0], 2)


if __name__ == "__main__": unittest.main()
